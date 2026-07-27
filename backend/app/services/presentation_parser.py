from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation


@dataclass(frozen=True)
class ExtractedSlide:
    slide_number: int
    title: str
    text: str


def _clean_text(value: str) -> str:
    lines = [
        line.strip()
        for line in value.splitlines()
        if line.strip()
    ]
    return "\n".join(lines)


def extract_presentation(file_path: Path) -> list[ExtractedSlide]:
    """Extract readable text and tables from a PPTX file."""
    if file_path.suffix.lower() != ".pptx":
        raise ValueError("Only PPTX presentation files are supported.")

    presentation = Presentation(str(file_path))
    extracted_slides: list[ExtractedSlide] = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = _clean_text(title_shape.text or "") if title_shape else ""
        content_parts: list[str] = []

        for shape in slide.shapes:
            if shape is title_shape:
                continue

            if getattr(shape, "has_text_frame", False):
                shape_text = _clean_text(shape.text or "")
                if shape_text:
                    content_parts.append(shape_text)

            if getattr(shape, "has_table", False):
                rows: list[str] = []
                for row in shape.table.rows:
                    cells = [_clean_text(cell.text or "") for cell in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    content_parts.append("\n".join(rows))

        slide_text = "\n\n".join(content_parts).strip()
        if not title and not slide_text:
            continue

        extracted_slides.append(
            ExtractedSlide(
                slide_number=slide_number,
                title=title or f"Slide {slide_number}",
                text=slide_text,
            )
        )

    if not extracted_slides:
        raise ValueError("No extractable text was found in the presentation.")

    return extracted_slides
